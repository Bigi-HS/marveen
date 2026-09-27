#!/usr/bin/env python3
"""Kahoot finalize: answer-position balancing + colored xlsx + PDF cheat-sheet.

Usage:
    python3 scripts/kahoot-finalize.py INPUT.xlsx [OUTPUT.pdf [TITLE]]

If OUTPUT.pdf is omitted it is derived from INPUT (same name, .pdf ext).
When chained from kahoot-build.py --finalize the xlsx is rewritten IN PLACE
(balanced correct positions). In standalone use the same applies -- make sure
you intend to overwrite the source file or point to a copy.
"""

import sys
from collections import Counter
from pathlib import Path


def _find_font(name: str, bold: bool = False) -> str:
    """Return a font path that exists on this host, falling back to a default."""
    suffix = "-Bold" if bold else ""
    candidates = [
        f"/usr/share/fonts/truetype/dejavu/DejaVuSans{suffix}.ttf",
        f"/usr/share/fonts/dejavu/DejaVuSans{suffix}.ttf",
        f"/usr/share/fonts/DejaVuSans{suffix}.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return p
    # reportlab will fall back to Helvetica if registration fails -- not ideal
    # but better than a hard crash on systems without DejaVu
    return candidates[0]


def main() -> int:
    src_arg = sys.argv[1] if len(sys.argv) > 1 else None
    if src_arg is None:
        print(__doc__)
        return 2

    src = Path(src_arg)
    if not src.exists():
        print(f"HIBA: a forras nem letezik: {src}")
        return 1

    pdf_path = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".pdf")
    title = sys.argv[3] if len(sys.argv) > 3 else "Kahoot Kvíz"

    # --- load xlsx ----------------------------------------------------------
    from openpyxl import load_workbook
    wb = load_workbook(src)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    data = []
    for idx, r in enumerate(rows[1:], start=2):
        if not r[0]:
            continue
        q = str(r[0])
        ans = [r[1], r[2], r[3], r[4]]
        t = r[5]
        try:
            ci = int(str(r[6]).split(",")[0].split(";")[0].strip()) - 1
        except (ValueError, AttributeError):
            print(f"HIBA: sor {idx}: nem ertelmezheto helyes-index: {r[6]!r}")
            return 1
        present = [i for i in range(4) if ans[i] is not None]
        if ci < 0 or ci >= 4 or ci not in present:
            print(f"HIBA: sor {idx}: helyes-index ({ci+1}) kiesik az ervenyes valaszok korul: {present}")
            return 1
        data.append([q, ans, t, ans[ci]])

    if not data:
        print("HIBA: nincsenek adatsorok az xlsx-ben")
        return 1

    # --- balanced round-robin correct position ------------------------------
    rr = [0, 1, 2, 3]
    k = 0
    out = []
    for q, ans, t, ctext in data:
        present = [i for i in range(4) if ans[i] is not None]
        others = [ans[i] for i in present if ans[i] != ctext]
        target = rr[k % 4]
        while target not in present:
            k += 1
            target = rr[k % 4]
        k += 1
        newans = [None, None, None, None]
        newans[target] = ctext
        oi = 0
        for p in present:
            if p == target:
                continue
            newans[p] = others[oi]
            oi += 1
        out.append((q, newans, t, target))
    data = out
    print("New correct positions:", Counter(d[3] + 1 for d in data))

    # --- rewrite xlsx in place ----------------------------------------------
    from openpyxl.styles import PatternFill, Font
    green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    gbold = Font(color="006100", bold=True)
    for i, (q, ans, t, ci) in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=q)
        for j in range(4):
            c = ws.cell(row=i, column=j + 2, value=ans[j])
            c.fill = green if j == ci else PatternFill(fill_type=None)
            c.font = gbold if j == ci else Font(color="000000")
        ws.cell(row=i, column=6, value=t)
        ws.cell(row=i, column=7, value=str(ci + 1))
    wb.save(src)
    print("XLSX saved:", src)

    # --- PDF cheat-sheet ----------------------------------------------------
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph

    try:
        pdfmetrics.registerFont(TTFont("DejaVu", _find_font("DejaVu")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", _find_font("DejaVu", bold=True)))
        font_regular, font_bold = "DejaVu", "DejaVu-Bold"
    except Exception as exc:
        print(f"  figyelem: DejaVu regisztracio sikertelen ({exc}), Helvetica-ra allunk vissza")
        font_regular, font_bold = "Helvetica", "Helvetica-Bold"

    qs = ParagraphStyle("q", fontName=font_bold, fontSize=10, leading=13,
                        textColor=colors.HexColor("#1a1a2e"))
    astyle = ParagraphStyle("a", fontName=font_regular, fontSize=9, leading=12)
    acor = ParagraphStyle("ac", fontName=font_bold, fontSize=9, leading=12,
                          textColor=colors.black)
    title_style = ParagraphStyle("t", fontName=font_bold, fontSize=18, leading=22,
                                 textColor=colors.HexColor("#0b6e4f"))
    sub = ParagraphStyle("s", fontName=font_regular, fontSize=10, leading=14,
                         textColor=colors.HexColor("#555555"))

    doc = SimpleDocTemplate(str(pdf_path), pagesize=A4,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm)
    el = [
        Paragraph(title, title_style),
        Paragraph(
            f"{len(data)} kérdés · a helyes válasz zöld háttérrel kiemelve · Kahoot import-kész",
            sub,
        ),
        Spacer(1, 6 * mm),
    ]
    gbg = colors.HexColor("#C6EFCE")
    letters = ["A", "B", "C", "D"]
    for n, (q, ans, t, ci) in enumerate(data, 1):
        tr = [[Paragraph(f"{n}. {q}", qs), ""]]
        cs = [
            ("SPAN", (0, 0), (1, 0)),
            ("FONTNAME", (0, 0), (-1, -1), font_regular),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("LINEBELOW", (0, 0), (-1, 0), 0.4, colors.HexColor("#cccccc")),
        ]
        ri = 1
        for j in range(4):
            if ans[j] is None:
                continue
            isc = j == ci
            tr.append([Paragraph(f"{letters[j]}.  {ans[j]}", acor if isc else astyle), ""])
            cs.append(("SPAN", (0, ri), (1, ri)))
            if isc:
                cs.append(("BACKGROUND", (0, ri), (1, ri), gbg))
            ri += 1
        tb = Table(tr, colWidths=[150 * mm, 20 * mm])
        tb.setStyle(TableStyle(cs))
        el.append(tb)
        el.append(Spacer(1, 3 * mm))
    doc.build(el)
    print("PDF saved:", pdf_path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
